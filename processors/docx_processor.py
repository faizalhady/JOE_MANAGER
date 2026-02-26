"""
DOCX Document Processor
Extracts text from Word documents, splitting by headings to preserve structure.
"""

import logging
from typing import List, Dict, Any
from pathlib import Path
from .base import BaseProcessor

logger = logging.getLogger(__name__)


class DOCXProcessor(BaseProcessor):

    supported_extensions = [".docx"]

    def extract(self, file_path: str) -> List[Dict[str, Any]]:
        """Extract text from DOCX, splitting by heading structure."""
        self.validate_file(file_path)
        from docx import Document

        file_name = Path(file_path).name
        doc = Document(file_path)

        sections = []
        current_section = {
            "heading": "Document Start",
            "heading_level": 0,
            "paragraphs": [],
        }

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Check if this paragraph is a heading
            if para.style.name.startswith("Heading"):
                # Save the previous section if it has content
                if current_section["paragraphs"]:
                    sections.append(self._section_to_dict(
                        current_section, file_name, len(sections)
                    ))

                # Start a new section
                level = self._get_heading_level(para.style.name)
                current_section = {
                    "heading": text,
                    "heading_level": level,
                    "paragraphs": [],
                }
            else:
                current_section["paragraphs"].append(text)

        # Don't forget the last section
        if current_section["paragraphs"]:
            sections.append(self._section_to_dict(
                current_section, file_name, len(sections)
            ))

        # Also extract tables
        for t_idx, table in enumerate(doc.tables):
            table_text = self._table_to_text(table)
            if table_text.strip():
                sections.append({
                    "content": table_text,
                    "metadata": {
                        "source": file_name,
                        "table_index": t_idx,
                    },
                    "chunk_type": "table",
                })

        logger.info(f"Extracted {len(sections)} sections from {file_name}")
        return sections

    def _section_to_dict(self, section: dict, file_name: str, index: int) -> Dict[str, Any]:
        """Convert a collected section into the standard output format."""
        content = "\n".join(section["paragraphs"])

        # Prepend the heading for context
        if section["heading"] and section["heading"] != "Document Start":
            content = f"{section['heading']}\n{content}"

        return {
            "content": content,
            "metadata": {
                "source": file_name,
                "heading": section["heading"],
                "heading_level": section["heading_level"],
                "section_index": index,
            },
            "chunk_type": "text",
        }

    def _get_heading_level(self, style_name: str) -> int:
        """Extract heading level number from style name."""
        # "Heading 1" → 1, "Heading 2" → 2, etc.
        try:
            return int(style_name.split()[-1])
        except (ValueError, IndexError):
            return 0

    def _table_to_text(self, table) -> str:
        """Convert a docx table to readable text."""
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        return "\n".join(rows)
