"""
PPTX (PowerPoint) Document Processor
Extracts text from slides, speaker notes, and tables.
"""

import logging
from typing import List, Dict, Any
from pathlib import Path
from .base import BaseProcessor

logger = logging.getLogger(__name__)


class PPTXProcessor(BaseProcessor):

    supported_extensions = [".pptx"]

    def extract(self, file_path: str) -> List[Dict[str, Any]]:
        """Extract text from PowerPoint slides, notes, and tables."""
        self.validate_file(file_path)
        from pptx import Presentation

        file_name = Path(file_path).name
        prs = Presentation(file_path)
        sections = []

        total_slides = len(prs.slides)
        logger.info(f"Processing PPTX: {file_name} ({total_slides} slides)")

        for slide_num, slide in enumerate(prs.slides, 1):
            # ─── Slide text (from all shapes) ───
            slide_texts = []
            for shape in slide.shapes:
                # Text frames (titles, body text, text boxes)
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)

                # Tables
                if shape.has_table:
                    table_text = self._table_to_text(shape.table)
                    if table_text.strip():
                        sections.append({
                            "content": table_text,
                            "metadata": {
                                "source": file_name,
                                "slide": slide_num,
                                "total_slides": total_slides,
                            },
                            "chunk_type": "table",
                        })

            if slide_texts:
                content = "\n".join(slide_texts)
                sections.append({
                    "content": content,
                    "metadata": {
                        "source": file_name,
                        "slide": slide_num,
                        "total_slides": total_slides,
                    },
                    "chunk_type": "text",
                })

            # ─── Speaker notes ───
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    sections.append({
                        "content": notes_text,
                        "metadata": {
                            "source": file_name,
                            "slide": slide_num,
                            "total_slides": total_slides,
                            "content_type": "speaker_notes",
                        },
                        "chunk_type": "text",
                    })

        logger.info(f"Extracted {len(sections)} sections from {file_name}")
        return sections

    def _table_to_text(self, table) -> str:
        """Convert a pptx table to readable text."""
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        return "\n".join(rows)
